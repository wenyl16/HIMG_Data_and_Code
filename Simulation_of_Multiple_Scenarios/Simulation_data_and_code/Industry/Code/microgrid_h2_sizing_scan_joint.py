"""
PV capacity scan and hydrogen storage sizing with joint storage cost sensitivity.
================================================================================
Differences from microgrid_h2_sizing_scan.py:
  - Original: fixed battery cost, with independent price sensitivity for each hydrogen component (Elec/H2/FC).
  - Joint version: multiply all four CAPEX terms (BAT/ELEC/H2/FC) by the same storage_factor.
Scan: PV capacity x utilization target (60%-100%) x storage_factor [1.0/0.9/0.7/0.5].
Battery capacity remains fixed at 5000 kWh; only its cost is scaled. PV and grid costs are unchanged.
The MILP model and constraints match the original; only the cost calculation changes. Solver: Gurobi.
"""

import os
import openpyxl
import gurobipy as gp
from gurobipy import GRB
import time
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np

# ============================================================
# 1. System parameters (same as the original H2 model)
# ============================================================

PV_RATED_ORIGINAL = 1000
PV_CAPEX = 530
PV_LIFE = 25
PV_OM_RATE = 0.015

PV_SCAN_MIN = 2200
PV_SCAN_MAX = 75000
PV_SCAN_STEP = 500

P_BAT_MAX = 2500
E_BAT_MAX = 5000
ETA_CHARGE = 0.95
ETA_DISCHARGE = 0.95
SOC_BAT_MIN = 0.1
SOC_BAT_MAX = 0.9
SOC_BAT_INIT = 0.5
BAT_CAPEX = 250
BAT_LIFE = 8
BAT_OM_RATE = 0.01

ETA_ELEC = 0.0205
ELEC_CAPEX = 1000
ELEC_LIFE = 10
ELEC_OM_RATE = 0.01

SOC_H2_MIN = 0.05
SOC_H2_MAX = 0.98
SOC_H2_INIT = 0.5
H2_CAPEX = 400
H2_LIFE = 20
H2_OM_RATE = 0.01

ETA_FC = 18.3315
FC_CAPEX = 1300
FC_LIFE = 10
FC_OM_RATE = 0.01

# --- Joint storage cost scaling: one multiplier for BAT/ELEC/H2/FC CAPEX ---
STORAGE_PRICE_FACTORS = [1.0, 0.9, 0.7, 0.5]

PRICE_VALLEY = 0.451 / 7
PRICE_FLAT = 0.738 / 7
PRICE_PEAK = 1.144 / 7
PRICE_SHARP = 1.24 / 7
PRICE_DEEP_VALLEY = 0.355 / 7

DISCOUNT_RATE = 0.05

CONSUMPTION_TARGETS = [0.6, 0.7, 0.8, 0.9, 1.0]
CONSUMPTION_TOL = 0.0005

T = 8760


# ============================================================
# 2. Helper functions (same as the original H2 model)
# ============================================================

def crf(rate, years):
    if rate == 0:
        return 1.0 / years
    return rate * (1 + rate) ** years / ((1 + rate) ** years - 1)


def load_raw_data(filepath=None):
    if filepath is None:
        filepath = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Annual_8760.xlsx")
    wb = openpyxl.load_workbook(filepath, read_only=True)
    pv_raw = [row[1] for row in wb["PV"].iter_rows(min_row=1, max_row=T, max_col=2, values_only=True)]
    load = [row[1] for row in wb["Load"].iter_rows(min_row=1, max_row=T, max_col=2, values_only=True)]
    wb.close()
    return pv_raw, load


_DAYS = [31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
_MONTH_START = [0]
for _d in _DAYS:
    _MONTH_START.append(_MONTH_START[-1] + _d * 24)
_SPRING_AUTUMN = {2, 3, 4, 5, 6, 9, 10, 11}


def _month_of(hour_idx):
    for _m in range(12):
        if _MONTH_START[_m] <= hour_idx < _MONTH_START[_m + 1]:
            return _m + 1
    return 12


def build_prices():
    prices = []
    for i in range(T):
        h = i % 24
        if _month_of(i) in _SPRING_AUTUMN:
            if 16 <= h < 23:
                p = PRICE_PEAK
            elif (7 <= h < 11) or (14 <= h < 16) or (h == 23):
                p = PRICE_FLAT
            else:
                p = PRICE_VALLEY
        else:
            if 18 <= h < 22:
                p = PRICE_SHARP
            elif (16 <= h < 18) or (22 <= h < 23):
                p = PRICE_PEAK
            elif (7 <= h < 11) or (14 <= h < 16) or (h == 23):
                p = PRICE_FLAT
            else:
                p = PRICE_VALLEY
        prices.append(p)
    return prices


def scale_pv(pv_raw, pv_rated):
    s = pv_rated / PV_RATED_ORIGINAL
    return [v * s for v in pv_raw]


# ============================================================
# 3. Solve one MILP scenario (original H2 model, with storage_factor applied to costs)
# ============================================================

def solve_scenario(pv, load, prices, pv_rated, consumption_target, storage_factor=1.0):
    """
    storage_factor: shared CAPEX multiplier for BAT/ELEC/H2/FC (default 1.0 = original costs).
    """
    total_pv = sum(pv)
    max_pv = max(pv)
    max_load = max(load)
    total_load = sum(load)

    if total_pv == 0:
        return None
    if total_load == 0:
        return None

    surplus_list = [max(pv[t] - load[t], 0) for t in range(T)]
    max_surplus = max(surplus_list)
    total_surplus = sum(surplus_list)

    elec_cap_ub = max(max_surplus, 1.0)
    fc_cap_ub = max_load
    h2_tank_ub = max(1000, total_surplus * ETA_ELEC * 0.5)

    model = gp.Model("h2_sizing_joint")
    model.setParam("OutputFlag", 0)
    model.setParam("TimeLimit", 600)
    mipgap = 0.005 if consumption_target < 1.0 else 0.02
    model.setParam("MIPGap", mipgap)
    model.setParam("Threads", 1)

    idx = range(T)

    P_ELEC_CAP = model.addVar(lb=0, ub=elec_cap_ub, name="P_ELEC_CAP")
    P_FC_CAP = model.addVar(lb=0, ub=fc_cap_ub, name="P_FC_CAP")
    H2_TANK = model.addVar(lb=0, ub=h2_tank_ub, name="H2_TANK")

    h2_on = model.addVar(vtype=GRB.BINARY, name="h2_on")
    model.addConstr(P_ELEC_CAP <= elec_cap_ub * h2_on, name="elec_link")
    model.addConstr(P_FC_CAP >= max_load * h2_on, name="fc_min_link")
    model.addConstr(P_FC_CAP <= fc_cap_ub * h2_on, name="fc_max_link")
    model.addConstr(H2_TANK <= h2_tank_ub * h2_on, name="tank_link")

    P_grid = model.addVars(idx, lb=0, name="P_grid")
    P_PV_load = model.addVars(idx, lb=0, name="P_PV_load")
    P_PV_batt = model.addVars(idx, lb=0, name="P_PV_batt")
    P_PV_H2 = model.addVars(idx, lb=0, name="P_PV_H2")
    P_PV_curt = model.addVars(idx, lb=0, name="P_PV_curt")

    P_bat_ch = model.addVars(idx, lb=0, ub=P_BAT_MAX, name="P_bat_ch")
    P_bat_dis = model.addVars(idx, lb=0, ub=P_BAT_MAX, name="P_bat_dis")
    SOC_bat = model.addVars(idx, lb=SOC_BAT_MIN, ub=SOC_BAT_MAX, name="SOC_bat")

    P_elec = model.addVars(idx, lb=0, name="P_elec")
    P_FC = model.addVars(idx, lb=0, name="P_FC")
    m_h2 = model.addVars(idx, lb=0, name="m_h2")

    u_ch = model.addVars(idx, vtype=GRB.BINARY, name="u_ch")
    u_dis = model.addVars(idx, vtype=GRB.BINARY, name="u_dis")

    crf_pv = crf(DISCOUNT_RATE, PV_LIFE)
    crf_bat = crf(DISCOUNT_RATE, BAT_LIFE)
    crf_elec = crf(DISCOUNT_RATE, ELEC_LIFE)
    crf_h2 = crf(DISCOUNT_RATE, H2_LIFE)
    crf_fc = crf(DISCOUNT_RATE, FC_LIFE)

    # Scale all four storage CAPEX terms by storage_factor; retain PV and grid costs.
    cost_pv = (crf_pv + PV_OM_RATE) * PV_CAPEX * pv_rated
    cost_bat = (crf_bat + BAT_OM_RATE) * (BAT_CAPEX * storage_factor) * E_BAT_MAX
    cost_elec = (crf_elec + ELEC_OM_RATE) * (ELEC_CAPEX * storage_factor) * P_ELEC_CAP
    cost_h2 = (crf_h2 + H2_OM_RATE) * (H2_CAPEX * storage_factor) * H2_TANK
    cost_fc = (crf_fc + FC_OM_RATE) * (FC_CAPEX * storage_factor) * P_FC_CAP
    cost_grid = gp.quicksum(prices[t] * P_grid[t] for t in idx)

    model.setObjective(cost_pv + cost_bat + cost_elec + cost_h2 + cost_fc + cost_grid, GRB.MINIMIZE)

    for t in idx:
        model.addConstr(P_PV_load[t] + P_PV_batt[t] + P_PV_H2[t] + P_PV_curt[t] == pv[t], name=f"pv_alloc_{t}")
        model.addConstr(P_elec[t] == P_PV_H2[t], name=f"elec_pv_{t}")
        model.addConstr(P_bat_ch[t] == P_PV_batt[t], name=f"bat_pv_{t}")
        model.addConstr(P_elec[t] <= P_ELEC_CAP, name=f"elec_cap_{t}")
        model.addConstr(P_FC[t] <= P_FC_CAP, name=f"fc_cap_{t}")
        model.addConstr(P_bat_ch[t] <= pv[t] * u_ch[t], name=f"ch_flag_{t}")
        model.addConstr(P_bat_dis[t] <= load[t] * u_dis[t], name=f"dis_flag_{t}")
        model.addConstr(u_ch[t] + u_dis[t] <= 1, name=f"excl_{t}")
        if t == 0:
            model.addConstr(SOC_bat[t] == SOC_BAT_INIT + (ETA_CHARGE * P_bat_ch[t] - P_bat_dis[t] / ETA_DISCHARGE) / E_BAT_MAX, name=f"soc_bat_{t}")
        else:
            model.addConstr(SOC_bat[t] == SOC_bat[t - 1] + (ETA_CHARGE * P_bat_ch[t] - P_bat_dis[t] / ETA_DISCHARGE) / E_BAT_MAX, name=f"soc_bat_{t}")
        if t == 0:
            model.addConstr(m_h2[t] == SOC_H2_INIT * H2_TANK + ETA_ELEC * P_elec[t] - P_FC[t] / ETA_FC, name=f"h2_{t}")
        else:
            model.addConstr(m_h2[t] == m_h2[t - 1] + ETA_ELEC * P_elec[t] - P_FC[t] / ETA_FC, name=f"h2_{t}")
        model.addConstr(m_h2[t] >= SOC_H2_MIN * H2_TANK, name=f"h2_lo_{t}")
        model.addConstr(m_h2[t] <= SOC_H2_MAX * H2_TANK, name=f"h2_hi_{t}")
        model.addConstr(load[t] == P_PV_load[t] + P_bat_dis[t] + P_FC[t] + P_grid[t], name=f"balance_{t}")

    model.addConstr(SOC_bat[T - 1] == SOC_BAT_INIT, name="bat_final")
    model.addConstr(m_h2[T - 1] == SOC_H2_INIT * H2_TANK, name="h2_final")

    curtail_sum = gp.quicksum(P_PV_curt[t] for t in idx)
    cons_tol = CONSUMPTION_TOL
    max_curtail = (1 - consumption_target + cons_tol) * total_pv
    min_curtail = max(0, (1 - consumption_target - cons_tol) * total_pv)
    model.addConstr(curtail_sum <= max_curtail, name="curt_hi")
    if consumption_target < 1.0:
        model.addConstr(curtail_sum >= min_curtail, name="curt_lo")

    model.optimize()

    if model.status not in (GRB.OPTIMAL, GRB.TIME_LIMIT, GRB.SUBOPTIMAL):
        return None
    if model.SolCount == 0:
        return None

    elec_cap = P_ELEC_CAP.X
    fc_cap = P_FC_CAP.X
    h2_cap = H2_TANK.X

    total_grid = sum(P_grid[t].X for t in idx)
    total_grid_cost = sum(prices[t] * P_grid[t].X for t in idx)
    total_pv_load = sum(P_PV_load[t].X for t in idx)
    total_bat_dis = sum(P_bat_dis[t].X for t in idx)
    total_fc_gen = sum(P_FC[t].X for t in idx)
    total_curtail = sum(P_PV_curt[t].X for t in idx)

    energy_pen = (total_pv_load + total_bat_dis + total_fc_gen) / total_load * 100
    actual_cons = (total_pv - total_curtail) / total_pv * 100
    lcoe = model.ObjVal / total_load

    # Annualized investment breakdown after cost scaling, consistent with the objective function.
    ann_pv = (crf_pv + PV_OM_RATE) * PV_CAPEX * pv_rated
    ann_bat = (crf_bat + BAT_OM_RATE) * (BAT_CAPEX * storage_factor) * E_BAT_MAX
    ann_elec = (crf_elec + ELEC_OM_RATE) * (ELEC_CAPEX * storage_factor) * elec_cap
    ann_h2 = (crf_h2 + H2_OM_RATE) * (H2_CAPEX * storage_factor) * h2_cap
    ann_fc = (crf_fc + FC_OM_RATE) * (FC_CAPEX * storage_factor) * fc_cap

    return {
        "storage_factor": storage_factor,
        "PV_MW": pv_rated / 1000,
        "cons_target_pct": consumption_target * 100,
        "elec_cap_kW": round(elec_cap, 1),
        "fc_cap_kW": round(fc_cap, 1),
        "h2_tank_kg": round(h2_cap, 1),
        "h2_tank_MWh": round(round(h2_cap, 1) * 33.33 / 1000, 4),
        "grid_kWh": round(total_grid, 0),
        "grid_cost_USD": round(total_grid_cost, 0),
        "curtail_kWh": round(total_curtail, 0),
        "actual_cons_pct": round(actual_cons, 2),
        "energy_pen_pct": round(energy_pen, 2),
        "ann_pv": round(ann_pv, 0),
        "ann_bat": round(ann_bat, 0),
        "ann_elec": round(ann_elec, 0),
        "ann_h2": round(ann_h2, 0),
        "ann_fc": round(ann_fc, 0),
        "total_cost_USD": round(model.ObjVal, 0),
        "LCOE": round(lcoe, 4),
    }


# ============================================================
# 4. Save results
# ============================================================

def save_to_excel(results, filename=None):
    if filename is None:
        filename = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                "sensitivity_h2_sizing_joint_results.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"

    headers = [
        "storage_factor",
        "PV_MW", "cons_target_%",
        "elec_cap_kW", "fc_cap_kW", "h2_tank_kg", "h2_tank_MWh",
        "grid_kWh", "grid_cost_USD", "curtail_kWh",
        "actual_cons_%", "energy_pen_%",
        "ann_pv", "ann_bat", "ann_elec", "ann_h2", "ann_fc",
        "total_cost_USD", "LCOE_USD_per_kWh",
    ]
    ws.append(headers)

    for r in results:
        # Exclude cases with no hydrogen system (h2_on=0 and all three component capacities equal to zero).
        if r["elec_cap_kW"] == 0 and r["fc_cap_kW"] == 0 and r["h2_tank_kg"] == 0:
            continue
        ws.append([
            r["storage_factor"],
            r["PV_MW"], r["cons_target_pct"],
            r["elec_cap_kW"], r["fc_cap_kW"], r["h2_tank_kg"], r["h2_tank_MWh"],
            r["grid_kWh"], r["grid_cost_USD"], r["curtail_kWh"],
            r["actual_cons_pct"], r["energy_pen_pct"],
            r["ann_pv"], r["ann_bat"], r["ann_elec"], r["ann_h2"], r["ann_fc"],
            r["total_cost_USD"], r["LCOE"],
        ])

    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 2

    wb.save(filename)
    wb.close()
    print(f"Results saved to {filename}")


# ============================================================
# 5. Visualization by storage_factor, following the battery version's style
# ============================================================

def plot_results(results):
    colors = ["tab:blue", "tab:orange", "tab:green", "tab:red", "tab:purple", "tab:brown"]
    out_dir = os.path.dirname(os.path.abspath(__file__))
    factor_colors = {1.0: "tab:blue", 0.9: "tab:orange", 0.7: "tab:green", 0.5: "tab:red"}

    base = [r for r in results if r["storage_factor"] == 1.0]

    # Figure 1: LCOE vs energy penetration (baseline storage=1.0).
    fig1, ax1 = plt.subplots(figsize=(11, 7))
    for i, ct in enumerate(CONSUMPTION_TARGETS):
        ct_pct = ct * 100
        pts = sorted([(r["energy_pen_pct"], r["LCOE"]) for r in base if r["cons_target_pct"] == ct_pct])
        if pts:
            xs, ys = zip(*pts)
            ax1.plot(xs, ys, linewidth=1.5, color=colors[i], marker="o", markersize=2, label=f"{ct_pct:.0f}%")
    ax1.set_xlabel("Energy Penetration (%)")
    ax1.set_ylabel("LCOE (USD/kWh)")
    ax1.set_title("H2 Sizing Joint (baseline): LCOE vs Energy Penetration")
    ax1.legend(title="Consumption", loc="best", fontsize=9)
    ax1.grid(True, alpha=0.3)
    fig1.tight_layout()
    fig1.savefig(os.path.join(out_dir, "plot_joint_lcoe_vs_penetration.png"), dpi=150)
    plt.close(fig1)

    # Figure 2: LCOE vs PV capacity (baseline).
    fig2, ax2 = plt.subplots(figsize=(11, 7))
    for i, ct in enumerate(CONSUMPTION_TARGETS):
        ct_pct = ct * 100
        pts = sorted([(r["PV_MW"], r["LCOE"]) for r in base if r["cons_target_pct"] == ct_pct])
        if pts:
            xs, ys = zip(*pts)
            ax2.plot(xs, ys, linewidth=1.5, color=colors[i], label=f"{ct_pct:.0f}%")
    ax2.set_xlabel("PV Capacity (MW)")
    ax2.set_ylabel("LCOE (USD/kWh)")
    ax2.set_title("H2 Sizing Joint (baseline): LCOE vs PV Capacity")
    ax2.legend(title="Consumption", loc="best", fontsize=9)
    ax2.grid(True, alpha=0.3)
    fig2.tight_layout()
    fig2.savefig(os.path.join(out_dir, "plot_joint_lcoe_vs_pv.png"), dpi=150)
    plt.close(fig2)

    # Figure 3: each H2 component's capacity vs PV, with curves by storage_factor and one panel per ct.
    for cap_name, cap_key, ylab in [("Electrolyzer", "elec_cap_kW", "kW"),
                                    ("H2 Tank", "h2_tank_kg", "kg"),
                                    ("Fuel Cell", "fc_cap_kW", "kW")]:
        fig, sub = plt.subplots(2, 3, figsize=(16, 9))
        flat = sub.flatten()
        for i, ct in enumerate(CONSUMPTION_TARGETS):
            ax = flat[i]
            ct_pct = ct * 100
            for factor in STORAGE_PRICE_FACTORS:
                pts = sorted([(r["PV_MW"], r[cap_key]) for r in results
                              if r["cons_target_pct"] == ct_pct and r["storage_factor"] == factor])
                if pts:
                    xs, ys = zip(*pts)
                    ax.plot(xs, ys, linewidth=1.3, color=factor_colors[factor], label=f"x{factor}")
            ax.set_title(f"Consumption {ct_pct:.0f}%")
            ax.set_xlabel("PV Capacity (MW)")
            ax.set_ylabel(ylab)
            ax.legend(title="storage_factor", fontsize=8)
            ax.grid(True, alpha=0.3)
        flat[5].axis("off")
        fig.suptitle(f"Joint Storage Cost Sensitivity — Optimal {cap_name} ({ylab})", fontsize=13)
        fig.tight_layout()
        fig.savefig(os.path.join(out_dir, f"plot_joint_{cap_key.replace('_kW','').replace('_kg','')}_sensitivity.png"), dpi=150)
        plt.close(fig)

    # Figure 4: LCOE vs storage_factor (mean and min/max range).
    fig_lcoe, ax_l = plt.subplots(figsize=(10, 6))
    factors_sorted = sorted(STORAGE_PRICE_FACTORS)
    means, lo, hi = [], [], []
    for f in factors_sorted:
        vals = [r["LCOE"] for r in results if r["storage_factor"] == f]
        if vals:
            means.append(np.mean(vals)); lo.append(min(vals)); hi.append(max(vals))
        else:
            means.append(np.nan); lo.append(np.nan); hi.append(np.nan)
    ax_l.plot(factors_sorted, means, "o-", color="tab:blue", label="mean")
    ax_l.fill_between(factors_sorted, lo, hi, color="tab:blue", alpha=0.15, label="min-max")
    ax_l.set_title("LCOE Sensitivity to Joint Storage Cost Factor (over all PV & consumption)")
    ax_l.set_xlabel("storage_factor (battery + whole H2 system)")
    ax_l.set_ylabel("LCOE (USD/kWh)")
    ax_l.legend(fontsize=9)
    ax_l.grid(True, alpha=0.3)
    fig_lcoe.tight_layout()
    fig_lcoe.savefig(os.path.join(out_dir, "plot_joint_lcoe_sensitivity.png"), dpi=150)
    plt.close(fig_lcoe)

    print("Figures saved: plot_joint_lcoe_vs_penetration.png, plot_joint_lcoe_vs_pv.png, "
          "plot_joint_elec_cap_sensitivity.png, plot_joint_h2_tank_sensitivity.png, "
          "plot_joint_fc_cap_sensitivity.png, plot_joint_lcoe_sensitivity.png")


# ============================================================
# 6. Main program
# ============================================================

if __name__ == "__main__":
    print("Loading data...")
    pv_raw, load_data = load_raw_data()
    prices = build_prices()

    pv_values = list(range(PV_SCAN_MIN, PV_SCAN_MAX + PV_SCAN_STEP, PV_SCAN_STEP))
    n_pv = len(pv_values)
    n_ct = len(CONSUMPTION_TARGETS)
    n_sf = len(STORAGE_PRICE_FACTORS)
    n_total = n_pv * n_ct * n_sf

    all_results = []
    n_done = 0
    n_ok = 0
    t_start = time.time()

    print(f"Scan (joint storage cost scaling): {n_pv} PV × {n_ct} utilization targets × {n_sf} storage_factor = {n_total} scenarios")
    print(f"PV range: {PV_SCAN_MIN/1000:.1f}~{PV_SCAN_MAX/1000:.0f} MW, step {PV_SCAN_STEP/1000:.1f} MW")
    print(f"storage_factor (battery and all hydrogen components): {STORAGE_PRICE_FACTORS}")
    print()

    for storage_factor in STORAGE_PRICE_FACTORS:
        for pv_kw in pv_values:
            pv = scale_pv(pv_raw, pv_kw)
            for ct in CONSUMPTION_TARGETS:
                n_done += 1

                if n_done % 200 == 0 or n_done == n_total:
                    elapsed = time.time() - t_start
                    rate = n_done / elapsed if elapsed > 0 else 0
                    eta = (n_total - n_done) / rate if rate > 0 else 0
                    print(f"  Progress {n_done}/{n_total} ({n_done/n_total*100:.0f}%)  "
                          f"Feasible {n_ok}  Elapsed {elapsed:.0f}s  Remaining ~{eta:.0f}s")

                r = solve_scenario(pv, load_data, prices, pv_kw, ct, storage_factor=storage_factor)
                if r is not None:
                    all_results.append(r)
                    n_ok += 1

    elapsed = time.time() - t_start
    print(f"\nScan complete: {n_ok}/{n_total} feasible, elapsed {elapsed:.1f}s")

    if all_results:
        all_results.sort(key=lambda r: (r["storage_factor"], r["PV_MW"], r["cons_target_pct"]))
        save_to_excel(all_results)
        plot_results(all_results)
    else:
        print("No feasible solution; check the parameter settings")
