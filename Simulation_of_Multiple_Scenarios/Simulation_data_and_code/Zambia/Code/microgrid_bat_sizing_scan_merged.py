"""
Zambia mining PV-battery microgrid sizing: merged/brownfield version retaining the existing battery
======================================================================================
Differences from greenfield microgrid_bat_sizing_scan.py:
  - E_BAT_MAX = 15000 kWh retains the existing 7500 kW/15000 kWh battery; E_TOTAL >= 15000 and power=E_TOTAL/2 (minimum 7500 kW).
  - BAT_PRICE_FACTORS = [1.0]; overseas scenarios use only the baseline factor.
  - Apply full CAPEX to all E_TOTAL, including the existing-capacity floor; existing capacity is repriced, not treated as a sunk cost.
Data: Annual_8760.xlsx; PV and Load sheets, each with hour-index and kW columns, no header, and 8760 hours.
System: scanned PV capacity referenced to a 1 MW nameplate rating, plus one battery (E_TOTAL decision variable; power=E/2).
Electricity price: no price column; assume a flat Zambia CEC mining tariff of 0.15 USD/kWh, already in USD.
Utilization targets: 60%-100%. Objective: minimize annualized total cost. Solver: Gurobi.
Differences from Norway: no wind, two sheets without headers, and a flat USD tariff requiring no currency conversion.
"""

import os
import openpyxl
import gurobipy as gp
from gurobipy import GRB
import time
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from scipy.interpolate import PchipInterpolator

# ============================================================
# 1. System parameters
# ============================================================

# PV: 1 MW nameplate reference, measured peak 830.65 kW.
PV_RATED_ORIGINAL = 1000         # kW: nameplate rating; the measured 830.65 kW peak reflects approximately 17% losses.
PV_CAPEX = 658                   # $/kW (nameplate)
PV_LIFE = 25
PV_OM_RATE = 0.015

# PV capacity scan (nameplate kW)
PV_SCAN_MIN = 50                # Match greenfield: the old 9000 kW minimum is infeasible at ct>=80% (14.6 GWh generation vs 11.3 GWh load, a ratio of 1.29; ct=100% allows about 6956 kW). Start at 50; E_BAT_MAX=15000 retains the brownfield battery floor.
PV_SCAN_MAX = 13000              # Raised to reach nearly 100% penetration at ct=60%, requiring about 11.6 MW.
PV_SCAN_STEP = 50                # Finer step than the previous 250 kW to capture seasonal storage peaks near each ct boundary.
WALL_FINE_STEP = 10             # Boundary refinement step in kW, used near 100% penetration.
WALL_PEN_ENTER = 0.90           # Energy penetration threshold for boundary refinement, expressed as a fraction.


def boundary_fine_tasks(results):
    """Generate boundary refinement tasks [(factor, pv_kw, ct_fraction), ...] from coarse-scan results.
    For each (factor, ct), start at the first PV capacity with energy_pen>=WALL_PEN_ENTER and end one coarse step beyond the largest feasible PV capacity.
    Use WALL_FINE_STEP, skipping capacities already solved, to fill the boundary near 100% penetration."""
    existing = {(round(r["PV_MW"] * 1000), r["cons_target_pct"], r["bat_price_factor"])
                for r in results}
    tasks = []
    for factor in sorted({r["bat_price_factor"] for r in results}):
        for ct_pct in sorted({r["cons_target_pct"] for r in results}):
            pts = sorted([r for r in results
                          if r["cons_target_pct"] == ct_pct and r["bat_price_factor"] == factor],
                         key=lambda r: r["PV_MW"])
            hi = [r for r in pts if r["energy_pen_pct"] >= WALL_PEN_ENTER * 100]
            if len(pts) < 2 or not hi:
                continue
            pv_start = hi[0]["PV_MW"] * 1000
            pv_end = pts[-1]["PV_MW"] * 1000 + PV_SCAN_STEP  # Probe one coarse step beyond the feasible range to locate the boundary.
            p = pv_start
            while p <= pv_end:
                pi = round(p)
                if (pi, ct_pct, factor) not in existing:
                    tasks.append((factor, pi, ct_pct / 100.0))
                p += WALL_FINE_STEP
    return tasks

# Backup grid supply cost: flat Zambia CEC mining tariff in USD/kWh.
GRID_PRICE_USD = 0.15

# Single battery storage, following Norway.
E_BAT_MAX = 15000                # kWh: total capacity lower bound retaining Zambia's existing 7500 kW/15000 kWh battery.
BAT_POWER_RATIO = 2              # Energy-to-power ratio = 2:1; minimum power = E_BAT_MAX/2 = 7500 kW.
BAT_CAPEX = 300                  # $/kWh
BAT_LIFE = 8
BAT_OM_RATE = 0.01

BAT_PRICE_FACTORS = [1.0]  # Overseas scenarios use only cost factor 1.0.

ETA_CHARGE = 0.95
ETA_DISCHARGE = 0.95
SOC_BAT_MIN = 0.1
SOC_BAT_MAX = 0.9
SOC_BAT_INIT = 0.5

# Economic parameters
DISCOUNT_RATE = 0.05

# Renewable utilization targets
CONSUMPTION_TARGETS = [0.6, 0.7, 0.8, 0.9, 1.0]
CONSUMPTION_TOL = 0.0005       # 0.05% tolerance for all utilization targets

# ============================================================
# 2. Helper functions
# ============================================================

def crf(rate, years):
    if rate == 0:
        return 1.0 / years
    return rate * (1 + rate) ** years / ((1 + rate) ** years - 1)


def load_raw_data(filepath=None):
    """Read the second column of PV/Load in Annual_8760.xlsx without headers; use flat GRID_PRICE_USD."""
    if filepath is None:
        filepath = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Annual_8760.xlsx")
    wb = openpyxl.load_workbook(filepath, read_only=True)
    pv_rows = list(wb["PV"].iter_rows(values_only=True))
    load_rows = list(wb["Load"].iter_rows(values_only=True))
    wb.close()

    # No header to skip; col[0]=hour-index, col[1]=power in kW.
    pv_raw = [max(0.0, float(r[1])) for r in pv_rows if r[1] is not None]
    load = [float(r[1]) for r in load_rows if r[1] is not None]

    T = min(8760, len(pv_raw), len(load))
    pv_raw = pv_raw[:T]
    load = load[:T]
    prices = [GRID_PRICE_USD] * T     # Flat electricity tariff

    print(f"  PV: peak={max(pv_raw):.1f}kW, annual generation={sum(pv_raw):.0f}kWh (nameplate {PV_RATED_ORIGINAL}kW)")
    print(f"  Load: peak={max(load):.1f}kW, annual consumption={sum(load):.0f}kWh")
    print(f"  Grid price: {GRID_PRICE_USD} USD/kWh (flat)")
    print(f"  Data length: {T}h")
    return pv_raw, load, prices


def scale_pv(pv_raw, pv_rated):
    s = pv_rated / PV_RATED_ORIGINAL
    return [v * s for v in pv_raw]


# ============================================================
# 3. Solve one PV-battery MILP without wind power.
# ============================================================

def solve_scenario(pv, load, prices, pv_rated, consumption_target, bat_price_factor=1.0):
    T = len(pv)
    ren = list(pv)
    total_ren = sum(ren)
    total_load = sum(load)

    if total_ren == 0:
        return None

    surplus_list = [max(ren[t] - load[t], 0) for t in range(T)]
    max_surplus = max(surplus_list)
    total_surplus = sum(surplus_list)

    add_ub = max(1000, min(total_surplus * 0.5, max(ren) * T * 0.25))
    etotal_ub = E_BAT_MAX + add_ub

    model = gp.Model("zambia_pv_bat_sizing_merged")
    model.setParam("OutputFlag", 0)
    model.setParam("TimeLimit", 300)
    mipgap = 0.005 if consumption_target < 0.8 else 0.02
    model.setParam("MIPGap", mipgap)
    model.setParam("Threads", 1)   # Use one solver thread, matching the parallel version for consistent serial/parallel runs.
    model.setParam("MIPFocus", 1)  # Prioritize feasibility to find solutions faster and reduce missing high-ct/high-PV points.

    idx = range(T)

    E_TOTAL = model.addVar(lb=E_BAT_MAX, ub=etotal_ub, name="E_TOTAL")

    P_grid = model.addVars(idx, lb=0, name="P_grid")
    P_ren_load = model.addVars(idx, lb=0, name="P_ren_load")
    P_ren_curt = model.addVars(idx, lb=0, name="P_ren_curt")

    P_bat_ch = model.addVars(idx, lb=0, name="P_bat_ch")
    P_bat_dis = model.addVars(idx, lb=0, name="P_bat_dis")
    E_bat = model.addVars(idx, lb=0, name="E_bat")

    u_ch = model.addVars(idx, vtype=GRB.BINARY, name="u_ch")
    u_dis = model.addVars(idx, vtype=GRB.BINARY, name="u_dis")

    crf_pv = crf(DISCOUNT_RATE, PV_LIFE)
    crf_bat = crf(DISCOUNT_RATE, BAT_LIFE)

    eff_bat_capex = BAT_CAPEX * bat_price_factor
    cost_pv = (crf_pv + PV_OM_RATE) * PV_CAPEX * pv_rated
    cost_bat = (crf_bat + BAT_OM_RATE) * eff_bat_capex * E_TOTAL
    cost_grid = gp.quicksum(prices[t] * P_grid[t] for t in idx)

    model.setObjective(cost_pv + cost_bat + cost_grid, GRB.MINIMIZE)

    for t in idx:
        model.addConstr(P_bat_ch[t] + P_ren_load[t] + P_ren_curt[t] == ren[t], name=f"ren_alloc_{t}")
        model.addConstr(P_bat_ch[t] <= E_TOTAL / BAT_POWER_RATIO, name=f"ch_lim_{t}")
        model.addConstr(P_bat_dis[t] <= E_TOTAL / BAT_POWER_RATIO, name=f"dis_lim_{t}")
        model.addConstr(P_bat_ch[t] <= ren[t] * u_ch[t], name=f"ch_{t}")
        model.addConstr(P_bat_dis[t] <= load[t] * u_dis[t], name=f"dis_{t}")
        model.addConstr(u_ch[t] + u_dis[t] <= 1, name=f"excl_{t}")
        if t == 0:
            model.addConstr(E_bat[t] == SOC_BAT_INIT * E_TOTAL
                           + (ETA_CHARGE * P_bat_ch[t] - P_bat_dis[t] / ETA_DISCHARGE), name=f"e_{t}")
        else:
            model.addConstr(E_bat[t] == E_bat[t - 1]
                           + (ETA_CHARGE * P_bat_ch[t] - P_bat_dis[t] / ETA_DISCHARGE), name=f"e_{t}")
        model.addConstr(E_bat[t] >= SOC_BAT_MIN * E_TOTAL, name=f"e_lo_{t}")
        model.addConstr(E_bat[t] <= SOC_BAT_MAX * E_TOTAL, name=f"e_hi_{t}")
        model.addConstr(load[t] == P_ren_load[t] + P_bat_dis[t] + P_grid[t], name=f"balance_{t}")

    model.addConstr(E_bat[T - 1] == SOC_BAT_INIT * E_TOTAL, name="bat_final")

    curtail_sum = gp.quicksum(P_ren_curt[t] for t in idx)
    cons_tol = CONSUMPTION_TOL   # Use a 0.05% tolerance for every utilization target.
    max_curtail = (1 - consumption_target + cons_tol) * total_ren
    min_curtail = max(0, (1 - consumption_target - cons_tol) * total_ren)
    model.addConstr(curtail_sum <= max_curtail, name="curt_hi")
    if consumption_target < 1.0:
        model.addConstr(curtail_sum >= min_curtail, name="curt_lo")

    model.optimize()

    if model.status not in (GRB.OPTIMAL, GRB.TIME_LIMIT, GRB.SUBOPTIMAL):
        return None
    if model.SolCount == 0:
        return None

    e_total = E_TOTAL.X
    p_total = e_total / BAT_POWER_RATIO

    total_grid = sum(P_grid[t].X for t in idx)
    total_grid_cost = sum(prices[t] * P_grid[t].X for t in idx)
    total_ren_load = sum(P_ren_load[t].X for t in idx)
    total_bat_dis = sum(P_bat_dis[t].X for t in idx)
    total_curtail = sum(P_ren_curt[t].X for t in idx)

    energy_pen = (total_ren_load + total_bat_dis) / total_load * 100
    actual_cons = (total_ren - total_curtail) / total_ren * 100
    lcoe = model.ObjVal / total_load

    return {
        "bat_price_factor": bat_price_factor,
        "PV_MW": pv_rated / 1000,
        "cons_target_pct": consumption_target * 100,
        "total_MWh": round(e_total / 1000, 2),
        "storage_duration_h": round(e_total / (total_load / T), 2),
        "total_MW": round(p_total / 1000, 2),
        "grid_kWh": round(total_grid, 0),
        "grid_cost_USD": round(total_grid_cost, 0),
        "curtail_kWh": round(total_curtail, 0),
        "actual_cons_pct": round(actual_cons, 2),
        "energy_pen_pct": round(energy_pen, 2),
        "total_cost_USD": round(model.ObjVal, 0),
        "LCOE": round(lcoe, 4),
    }


# ============================================================
# 4. Save results.
# ============================================================

def save_to_excel(results, filename=None):
    if filename is None:
        filename = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                "sensitivity_bat_sizing_merged_results.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"

    headers = [
        "bat_price_factor",
        "PV_MW", "cons_target_%",
        "total_MWh", "storage_duration_h", "total_MW",
        "grid_kWh", "grid_cost_USD", "curtail_kWh",
        "actual_cons_%", "energy_pen_%",
        "total_cost_USD", "LCOE_USD_per_kWh",
    ]
    ws.append(headers)

    for r in results:
        ws.append([
            r["bat_price_factor"],
            r["PV_MW"], r["cons_target_pct"],
            r["total_MWh"], r["storage_duration_h"], r["total_MW"],
            r["grid_kWh"], r["grid_cost_USD"], r["curtail_kWh"],
            r["actual_cons_pct"], r["energy_pen_pct"],
            r["total_cost_USD"], r["LCOE"],
        ])

    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 2

    wb.save(filename)
    print(f"Results saved to {filename}")


# ============================================================
# 5. Visualization
# ============================================================

def plot_results(results, avg_load_mw=None):
    colors = ["tab:blue", "tab:orange", "tab:green", "tab:red", "tab:purple"]
    out_dir = os.path.dirname(os.path.abspath(__file__))

    # Figure 1: LCOE vs penetration
    fig1, ax1 = plt.subplots(figsize=(11, 7))
    for i, ct in enumerate(CONSUMPTION_TARGETS):
        ct_pct = ct * 100
        pts = sorted([(r["energy_pen_pct"], r["LCOE"]) for r in results if r["cons_target_pct"] == ct_pct])
        if pts:
            xs, ys = zip(*pts)
            ax1.plot(xs, ys, linewidth=1.5, color=colors[i], marker="o", markersize=2, label=f"{ct_pct:.0f}%")
    ax1.set_xlabel("Energy Penetration (%)")
    ax1.set_ylabel("LCOE (USD/kWh)")
    ax1.set_title("Zambia PV+Bat: LCOE vs Energy Penetration")
    ax1.legend(title="Consumption", loc="best", fontsize=9)
    ax1.grid(True, alpha=0.3)
    fig1.tight_layout()
    fig1.savefig(os.path.join(out_dir, "plot_lcoe_vs_penetration.png"), dpi=150)
    plt.close(fig1)

    # Figure 2: PV capacity vs LCOE
    fig2, ax2 = plt.subplots(figsize=(11, 7))
    for i, ct in enumerate(CONSUMPTION_TARGETS):
        ct_pct = ct * 100
        pts = sorted([(r["PV_MW"], r["LCOE"]) for r in results if r["cons_target_pct"] == ct_pct])
        if pts:
            xs, ys = zip(*pts)
            ax2.plot(xs, ys, linewidth=1.5, color=colors[i], marker="o", markersize=2, label=f"{ct_pct:.0f}%")
    ax2.set_xlabel("PV Capacity (MW)")
    ax2.set_ylabel("LCOE (USD/kWh)")
    ax2.set_title("Zambia PV+Bat: LCOE vs PV Capacity")
    ax2.legend(title="Consumption", loc="best", fontsize=9)
    ax2.grid(True, alpha=0.3)
    fig2.tight_layout()
    fig2.savefig(os.path.join(out_dir, "plot_lcoe_vs_pv.png"), dpi=150)
    plt.close(fig2)

    # Figure 3: PV capacity vs battery capacity
    fig3, ax3 = plt.subplots(figsize=(11, 7))
    for i, ct in enumerate(CONSUMPTION_TARGETS):
        ct_pct = ct * 100
        pts = sorted([(r["PV_MW"], r["total_MWh"]) for r in results if r["cons_target_pct"] == ct_pct])
        if pts:
            xs, ys = zip(*pts)
            ax3.plot(xs, ys, linewidth=1.5, color=colors[i], marker="o", markersize=2, label=f"{ct_pct:.0f}%")
    ax3.set_xlabel("PV Capacity (MW)")
    ax3.set_ylabel("Total Battery Capacity (MWh)")
    ax3.set_title("Zambia: Optimal Battery Capacity vs PV Capacity")
    ax3.legend(title="Consumption", loc="best", fontsize=9)
    ax3.grid(True, alpha=0.3)
    fig3.tight_layout()
    fig3.savefig(os.path.join(out_dir, "plot_bat_capacity_vs_pv.png"), dpi=150)
    plt.close(fig3)

    # Figure 4: Battery capacity vs penetration
    fig4, ax4 = plt.subplots(figsize=(11, 7))
    for i, ct in enumerate(CONSUMPTION_TARGETS):
        ct_pct = ct * 100
        pts = sorted([(r["energy_pen_pct"], r["total_MWh"]) for r in results if r["cons_target_pct"] == ct_pct])
        if pts:
            xs, ys = zip(*pts)
            ax4.plot(xs, ys, linewidth=1.5, color=colors[i], marker="o", markersize=3, label=f"{ct_pct:.0f}%")
    ax4.set_xlabel("Energy Penetration (%)")
    ax4.set_ylabel("Total Battery Capacity (MWh)")
    ax4.set_title("Zambia PV+Bat: Battery Capacity vs Energy Penetration")
    ax4.legend(title="Consumption", loc="best", fontsize=9)
    ax4.grid(True, alpha=0.3)
    fig4.tight_layout()
    fig4.savefig(os.path.join(out_dir, "plot_capacity_vs_penetration.png"), dpi=150)
    plt.close(fig4)

    saved = ["plot_lcoe_vs_penetration.png", "plot_lcoe_vs_pv.png",
             "plot_bat_capacity_vs_pv.png", "plot_capacity_vs_penetration.png"]

    # Figure 5: Storage duration vs penetration (E_MWh / mean load in MW = h); PCHIP curves without markers.
    if avg_load_mw and avg_load_mw > 0:
        fig5, ax5 = plt.subplots(figsize=(11, 7))
        for i, ct in enumerate(CONSUMPTION_TARGETS):
            ct_pct = ct * 100
            pts = sorted([(r["energy_pen_pct"], r["total_MWh"] / avg_load_mw)
                          for r in results if r["cons_target_pct"] == ct_pct])
            if not pts:
                continue
            agg = {}
            for x, y in pts:
                agg.setdefault(x, []).append(y)
            xs_arr = np.asarray(sorted(agg), dtype=float)
            ys_arr = np.asarray([sum(agg[x]) / len(agg[x]) for x in sorted(agg)], dtype=float)
            if len(xs_arr) >= 2:
                xx = np.linspace(xs_arr[0], xs_arr[-1], 300)
                yy = PchipInterpolator(xs_arr, ys_arr)(xx)
                ax5.plot(xx, yy, linewidth=2.0, color=colors[i], label=f"{ct_pct:.0f}%")
            else:
                ax5.plot(xs_arr, ys_arr, linewidth=2.0, color=colors[i], label=f"{ct_pct:.0f}%")
        ax5.set_xlabel("Energy Penetration (%)")
        ax5.set_ylabel("Storage Duration (h)   [E / avg load]")
        ax5.set_title(f"Zambia PV+Bat: Storage Duration vs Energy Penetration "
                      f"(avg load = {avg_load_mw * 1000:.1f} kW)")
        ax5.legend(title="Consumption", loc="best", fontsize=9)
        ax5.grid(True, alpha=0.3)
        ax5.set_ylim(0, 200)
        fig5.tight_layout()
        fig5.savefig(os.path.join(out_dir, "plot_duration_vs_penetration.png"), dpi=150)
        plt.close(fig5)
        saved.append("plot_duration_vs_penetration.png")

    print("Figures saved: " + ", ".join(saved))


# ============================================================
# 6. Main program
# ============================================================

if __name__ == "__main__":
    print("Reading Zambia data...")
    pv_raw, load_data, price_data = load_raw_data()
    T = len(pv_raw)
    avg_load_mw = sum(load_data) / len(load_data) / 1000   # Mean load in MW for converting storage capacity to duration.

    pv_values = list(range(PV_SCAN_MIN, PV_SCAN_MAX + PV_SCAN_STEP, PV_SCAN_STEP))
    n_pv = len(pv_values)
    n_ct = len(CONSUMPTION_TARGETS)
    n_pf = len(BAT_PRICE_FACTORS)
    n_total = n_pv * n_ct * n_pf

    all_results = []
    n_done = 0
    n_ok = 0
    t_start = time.time()

    print(f"Scan (Zambia PV + battery): {n_pv} PV capacity levels x {n_ct} utilization targets x {n_pf} cost factors = {n_total} scenarios")
    print(f"PV range: {PV_SCAN_MIN/1000:.2f}~{PV_SCAN_MAX/1000:.1f} MW, step {PV_SCAN_STEP/1000:.2f} MW")
    print()

    for factor in BAT_PRICE_FACTORS:
        for pv_kw in pv_values:
            pv = scale_pv(pv_raw, pv_kw)
            for ct in CONSUMPTION_TARGETS:
                n_done += 1

                if n_done % 20 == 0 or n_done == n_total:
                    elapsed = time.time() - t_start
                    rate = n_done / elapsed if elapsed > 0 else 0
                    eta = (n_total - n_done) / rate if rate > 0 else 0
                    print(f"  Progress {n_done}/{n_total} ({n_done/n_total*100:.0f}%)  "
                          f"feasible {n_ok}  elapsed {elapsed:.0f}s  remaining ~{eta:.0f}s")

                r = solve_scenario(pv, load_data, price_data, pv_kw, ct, bat_price_factor=factor)
                if r is not None:
                    all_results.append(r)
                    n_ok += 1

    elapsed = time.time() - t_start
    print(f"\nScan complete: {n_ok}/{n_total} feasible, elapsed {elapsed:.1f}s")

    # Boundary refinement, matching the parallel version: use WALL_FINE_STEP where energy_pen>=WALL_PEN_ENTER.
    fine_tasks = boundary_fine_tasks(all_results)
    n_fine_ok = 0
    if fine_tasks:
        print(f"\nBoundary refinement: {len(fine_tasks)} additional tasks (step {WALL_FINE_STEP} kW, "
              f"pen>={WALL_PEN_ENTER*100:.0f}%)")
        for factor, pv_kw, ct_frac in fine_tasks:
            pv = scale_pv(pv_raw, pv_kw)
            r = solve_scenario(pv, load_data, price_data, pv_kw, ct_frac, bat_price_factor=factor)
            if r is not None:
                all_results.append(r)
                n_fine_ok += 1
        print(f"Boundary refinement complete: feasible {n_fine_ok}/{len(fine_tasks)}")

    if all_results:
        save_to_excel(all_results)
        plot_results(all_results, avg_load_mw=avg_load_mw)
    else:
        print("No feasible solution; check the parameters")
